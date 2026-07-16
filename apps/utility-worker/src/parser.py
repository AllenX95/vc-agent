import csv
import contextlib
import io
import importlib.metadata
import json
import sys
import time
import uuid
import xml.etree.ElementTree as ET
from pathlib import Path


def version(package):
    return importlib.metadata.version(package)


def source(command, kind="document", **locator):
    return {
        "relativePath": command["material"]["relativePath"],
        "sourceHash": command["material"]["sourceHash"],
        "locator": {"kind": kind, **locator},
    }


def parse(command):
    path = Path(command["material"]["absolutePath"])
    ext = path.suffix.lower()
    started = time.perf_counter()
    blocks, units, warnings, recovery = [], [], [], []

    def add(kind, text=None, rows=None, ref=None, level=None, tokens=None):
        block = {"id": f"b{len(blocks) + 1}", "type": kind, "source": ref or source(command)}
        if text is not None: block["text"] = text
        if rows is not None: block["rows"] = [["" if value is None else str(value) for value in row] for row in rows]
        if level is not None: block["level"] = level
        if tokens is not None: block["tokens"] = tokens
        blocks.append(block)
        return block["id"]

    if ext in (".txt", ".md", ".markdown", ".json", ".yaml", ".yml", ".xml"):
        text = path.read_text(encoding="utf-8-sig")
        if ext == ".json": json.loads(text)
        if ext in (".yaml", ".yml"):
            import yaml
            yaml.safe_load(text)
        if ext == ".xml": ET.fromstring(text)
        ids = []
        for index, line in enumerate(text.splitlines(), 1):
            if not line.strip(): continue
            heading = ext in (".md", ".markdown") and line.lstrip().startswith("#")
            level = min(6, len(line) - len(line.lstrip("#"))) if heading else None
            value = line.lstrip("#").strip() if heading else line
            ids.append(add("heading" if heading else "paragraph", text=value, level=level, ref=source(command, "line", index=index)))
        units.append({"index": 1, "name": path.name, "blockIds": ids})
        structure_kind = "structured_text" if ext in (".json", ".yaml", ".yml", ".xml") else "document"
        parser_id, parser_version = {
            ".json": ("json", "1.0.0"), ".yaml": ("pyyaml", version("PyYAML")),
            ".yml": ("pyyaml", version("PyYAML")), ".xml": ("xml", "1.0.0")
        }.get(ext, ("text", "1.0.0"))
        runtime = f"python {sys.version_info.major}.{sys.version_info.minor}"
    elif ext == ".csv":
        with path.open("r", encoding="utf-8-sig", newline="") as handle: rows = list(csv.reader(handle))
        block_id = add("table", rows=rows, ref=source(command, "document", name=path.name))
        units.append({"index": 1, "name": path.name, "blockIds": [block_id]})
        structure_kind, parser_id, parser_version, runtime = "table", "csv", "1.0.0", f"python {sys.version_info.major}.{sys.version_info.minor}"
    elif ext == ".docx":
        from docx import Document
        document = Document(path)
        ids = []
        for index, paragraph in enumerate(document.paragraphs, 1):
            if not paragraph.text.strip(): continue
            style = paragraph.style.name.lower() if paragraph.style else ""
            is_heading = style.startswith("heading")
            level = int(style.split()[-1]) if is_heading and style.split()[-1].isdigit() else None
            ids.append(add("heading" if is_heading else "paragraph", text=paragraph.text, level=level, ref=source(command, "path", path=f"paragraph[{index}]")))
        for index, table in enumerate(document.tables, 1):
            ids.append(add("table", rows=[[cell.text for cell in row.cells] for row in table.rows], ref=source(command, "path", path=f"table[{index}]")))
        units.append({"index": 1, "name": path.name, "blockIds": ids})
        structure_kind, parser_id, parser_version, runtime = "document", "python-docx", version("python-docx"), f"python {sys.version_info.major}.{sys.version_info.minor}"
    elif ext == ".pptx":
        from pptx import Presentation
        presentation = Presentation(path)
        for slide_index, slide in enumerate(presentation.slides, 1):
            ids = []
            for shape_index, shape in enumerate(slide.shapes, 1):
                ref = source(command, "slide", index=slide_index, path=f"shape[{shape_index}]")
                if getattr(shape, "has_table", False):
                    ids.append(add("table", rows=[[cell.text for cell in row.cells] for row in shape.table.rows], ref=ref))
                elif getattr(shape, "has_text_frame", False) and shape.text.strip():
                    ids.append(add("heading" if shape == slide.shapes.title else "paragraph", text=shape.text, level=1 if shape == slide.shapes.title else None, ref=ref))
            try:
                notes = slide.notes_slide.notes_text_frame.text.strip()
                if notes: ids.append(add("note", text=notes, ref=source(command, "slide", index=slide_index, path="speaker-notes")))
            except Exception: pass
            units.append({"index": slide_index, "name": f"Slide {slide_index}", "blockIds": ids})
        structure_kind, parser_id, parser_version, runtime = "slides", "python-pptx", version("python-pptx"), f"python {sys.version_info.major}.{sys.version_info.minor}"
    elif ext == ".xlsx":
        from openpyxl import load_workbook
        workbook = load_workbook(path, read_only=True, data_only=True)
        for sheet_index, sheet in enumerate(workbook.worksheets, 1):
            rows = [list(row) for row in sheet.iter_rows(values_only=True)]
            block_id = add("table", rows=rows, ref=source(command, "sheet", index=sheet_index, name=sheet.title))
            units.append({"index": sheet_index, "name": sheet.title, "blockIds": [block_id]})
        workbook.close()
        structure_kind, parser_id, parser_version, runtime = "workbook", "openpyxl", version("openpyxl"), f"python {sys.version_info.major}.{sys.version_info.minor}"
    elif ext == ".pdf":
        import pymupdf
        document = pymupdf.open(path)
        for page_index, page in enumerate(document, 1):
            ids = []
            words = page.get_text("words")
            for native in page.get_text("blocks"):
                text = native[4].strip()
                if text:
                    tokens = [{"text": word[4], "geometry": [word[0], word[1], word[2], word[3]]} for word in words if native[0] <= (word[0] + word[2]) / 2 <= native[2] and native[1] <= (word[1] + word[3]) / 2 <= native[3]]
                    ids.append(add("paragraph", text=text, tokens=tokens, ref=source(command, "page", index=page_index, geometry=[native[0], native[1], native[2], native[3]])))
            try:
                with contextlib.redirect_stdout(io.StringIO()):
                    detected_tables = page.find_tables().tables
                for table_index, table in enumerate(detected_tables, 1):
                    ids.append(add("table", rows=table.extract(), ref=source(command, "page", index=page_index, path=f"table[{table_index}]", geometry=list(table.bbox))))
            except Exception:
                warnings.append({"code": "PDF_TABLE_DETECTION_FAILED", "severity": "warning", "message": "Native table detection failed; other page content was retained.", "source": source(command, "page", index=page_index)})
            for image_index, image in enumerate(page.get_images(full=True), 1):
                ids.append(add("image", text=f"Embedded image xref {image[0]}", ref=source(command, "page", index=page_index, path=f"image[{image_index}]")))
            joined = " ".join(blocks[int(block_id[1:]) - 1].get("text", "") for block_id in ids)
            printable = sum(character.isprintable() for character in joined) / max(1, len(joined))
            if len(joined.strip()) < 20 or printable < 0.85:
                ref = source(command, "page", index=page_index)
                code = "OCR_UNAVAILABLE"
                warnings.append({"code": code, "severity": "warning", "message": "Native text is missing or unreliable; no page recovery adapter is registered.", "source": ref})
                recovery.append({"source": ref, "reason": "missing_text" if not joined.strip() else "unreliable_text", "status": "unavailable"})
            units.append({"index": page_index, "name": f"Page {page_index}", "blockIds": ids})
        document.close()
        structure_kind, parser_id, parser_version, runtime = "pages", "pymupdf", pymupdf.__version__, f"python {sys.version_info.major}.{sys.version_info.minor}"
    else:
        raise ValueError(f"Unsupported material type: {ext}")

    duration = max(0, round((time.perf_counter() - started) * 1000))
    warning_codes = sorted({warning["code"] for warning in warnings})
    return {
        "schemaVersion": 1,
        "parseId": str(uuid.uuid4()),
        "material": {key: command["material"][key] for key in ("id", "projectId", "relativePath", "mediaType", "sourceHash")},
        "parser": {"id": parser_id, "version": parser_version, "runtime": runtime},
        "createdAt": time.strftime("%Y-%m-%dT%H:%M:%S", time.gmtime()) + ".000Z",
        "structure": {"kind": structure_kind, "unitCount": len(units), "units": units},
        "blocks": blocks,
        "warnings": warnings,
        "recoveryRequests": recovery,
        "provenance": {"localOnly": True, "stages": [{"id": parser_id, "version": parser_version, "durationMs": duration, "status": "warning" if warnings else "completed", "warningCodes": warning_codes}]},
    }


if __name__ == "__main__":
    try:
        print(json.dumps(parse(json.load(sys.stdin)), ensure_ascii=False))
    except Exception as error:
        print(f"{type(error).__name__}: {error}", file=sys.stderr)
        sys.exit(1)
